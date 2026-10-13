"""Excel error values (#VALUE!, #REF!, #DIV/0!…) are reported as what they are: a broken formula in the client's file.

Also in .xlsx, .xls and .ods files, where the spreadsheet library reads them as empty (0.19.0). Such a cell holds
no data. It used to fail with the field type's own, misleading reason ("too few digits (1)" for
#DIV/0! as a phone, "unrecognized word 'value'" as a date). The outcome is unchanged - a phone row is still a bad
row, an optional value is still left blank - only the reason, the mistake check and "Suggest fixes" differ.
"""

import io
import sys
from pathlib import Path

import openpyxl
import pandas as pd
import pytest

from converter.checks import run_checks
from converter.engine import convert
from converter.mapping import FieldMapping, Mapping, profile_columns
from converter.normalizers import SHEET_ERROR_REASON, normalize, sheet_error
from converter.reader import read_file
from converter.sheeterrors import _XLS_ERROR_RECORD, error_cells

CODES = ["#VALUE!", "#REF!", "#DIV/0!", "#NAME?", "#NUM!", "#NULL!", "#SPILL!", "#CALC!", "#ERROR!", "#GETTING_DATA",
         "Err:502", "#value!", " #REF! "]


@pytest.mark.parametrize("code", CODES)
@pytest.mark.parametrize("field", ["phone_number", "transactions_last", "birthday", "consent"])
def test_every_field_type_names_the_excel_error(real_schema, code, field):
    assert normalize(code, real_schema.field(field), {}, {}) == ("", SHEET_ERROR_REASON)


@pytest.mark.parametrize("value", ["#N/A", "#n/a"])
def test_na_is_still_empty(real_schema, value):
    """#N/A keeps meaning "no value" (also typed by people), so a phone row says the number is missing."""
    assert normalize(value, real_schema.field("phone_number"), {}, {}) == ("", None)


@pytest.mark.parametrize("value", ["#5 Main St", "#1 customer", "Value!", "#REF", "Err:50", "#VALUE! 5165550100"])
def test_other_values_are_not_excel_errors(value):
    assert sheet_error(value) is None


def test_import_reasons_and_summary(real_schema):
    """The rejects report shows the reason and the code; all codes share one line in the summary."""
    df = pd.DataFrame({"Phone": ["#DIV/0!", "#REF!", "5165550100"], "Visit": ["10/01/2026", "", "#VALUE!"]})
    m = Mapping(fields={"phone_number": FieldMapping(sources=["Phone"], transform="copy"),
                        "transactions_last": FieldMapping(sources=["Visit"], transform="copy")})
    res = convert(df, real_schema, m, row_numbers=[2, 3, 4])
    assert res.output["phone_number"].tolist() == ["5165550100"] and res.output["transactions_last"].tolist() == [""]
    assert res.rejects["Reason"].tolist() == [
        "phone_number: Excel error (broken formula in the client's file) ('#DIV/0!')",
        "phone_number: Excel error (broken formula in the client's file) ('#REF!')"]
    assert res.reason_counts == {"phone_number: Excel error (broken formula in the client's file)": 2}
    assert res.blanked_counts == {"transactions_last: Excel error (broken formula in the client's file)": 1}


def test_mistake_check_points_at_the_export_not_the_column(real_schema):
    """A column of Excel errors gets its own warning, not "this is probably the wrong column"."""
    dates = ["10/01/2026", "10/02/2026"] + ["#REF!"] * 8
    df = pd.DataFrame({"Phone": [f"516555{1000 + i}" for i in range(10)], "Visit": dates})
    m = Mapping(fields={"phone_number": FieldMapping(sources=["Phone"], transform="copy"),
                        "transactions_last": FieldMapping(sources=["Visit"], transform="copy")})
    res = convert(df, real_schema, m)
    issues = [f.issue for f in run_checks(real_schema, m, df, res.audit, profile_columns(df))
              if f.field == "transactions_last"]
    assert any(i.startswith("8 value(s) from 'Visit' are Excel errors (e.g. '#REF!')") for i in issues)
    assert not any("can't be read" in i for i in issues)


# ---------------------------------------------------------------- spreadsheets (0.19.0)

FIXTURES = Path(__file__).parent / "fixtures"


def xlsx_with_errors() -> bytes:
    """The same table as fixtures/excel_errors.xls and .ods (made from it by LibreOffice), as an .xlsx."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Customers"
    for row in (["Phone", "Last Visit", "Birthday", "Notes"], ["#DIV/0!", "10/01/2026", "09/12/1985", "ok"],
                ["5165550101", "#VALUE!", "#N/A", "ok"], ["5165550102", "10/03/2026", "#REF!", "ok"]):
        ws.append(row)
    for ref in ("A2", "B3", "C3", "C4"):
        ws[ref].data_type = "e"  # error cells, as Excel stores a broken formula's result
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


SPREADSHEETS = {".xlsx": xlsx_with_errors, ".xls": lambda: (FIXTURES / "excel_errors.xls").read_bytes(),
                ".ods": lambda: (FIXTURES / "excel_errors.ods").read_bytes()}


@pytest.mark.parametrize("ext", SPREADSHEETS)
def test_spreadsheet_error_cells_keep_their_code(real_schema, ext):
    """The spreadsheet library reads error cells as empty: a phone said "required value is missing" and other
    fields were silently blank. The codes are put back, so .xlsx/.xls/.ods files report them like a CSV."""
    rr = read_file("f" + ext, SPREADSHEETS[ext]())
    assert rr.df.values.tolist() == [["#DIV/0!", "10/01/2026", "09/12/1985", "ok"],
                                     ["5165550101", "#VALUE!", "#N/A", "ok"],
                                     ["5165550102", "10/03/2026", "#REF!", "ok"]]
    m = Mapping(fields={"phone_number": FieldMapping(sources=["Phone"], transform="copy"),
                        "transactions_last": FieldMapping(sources=["Last Visit"], transform="copy"),
                        "birthday": FieldMapping(sources=["Birthday"], transform="copy")})
    res = convert(rr.df, real_schema, m)
    assert res.imported == 2 and res.reason_counts == {f"phone_number: {SHEET_ERROR_REASON}": 1}
    assert res.blanked_counts == {f"transactions_last: {SHEET_ERROR_REASON}": 1, f"birthday: {SHEET_ERROR_REASON}": 1}
    assert res.output["birthday"].tolist() == ["", ""]  # #N/A is empty, #REF! is an error: both blank


def test_error_cell_outside_the_values_is_added():
    """An error cell in a row and column with nothing else still shows (the grid is extended to reach it)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Phone", "Notes"])
    ws.append(["5165550100", "ok"])
    ws["C3"] = "#NAME?"
    ws["C3"].data_type = "e"
    buf = io.BytesIO()
    wb.save(buf)
    assert read_file("f.xlsx", buf.getvalue()).raw_preview == [["Phone", "Notes", ""], ["5165550100", "ok", ""],
                                                                ["", "", "#NAME?"]]


def test_xls_without_xlrd_reads_as_before(monkeypatch):
    """xlrd is optional: without it, .xls error cells stay empty (as before 0.19.0) and the file still reads."""
    monkeypatch.setitem(sys.modules, "xlrd", None)
    rr = read_file("f.xls", SPREADSHEETS[".xls"]())
    assert rr.df["Phone"].tolist() == ["", "5165550101", "5165550102"]


def test_xls_error_records_are_found_before_reading_again():
    """The quick byte search that decides whether an .xls is read a second time (for its error cells)."""
    boolerr = b"\x05\x02\x08\x00" + b"\x01\x00\x02\x00\x0f\x00" + b"\x07\x01"     # row 1, col 2: #DIV/0!
    formula = b"\x06\x00\x20\x00" + b"\x01\x00\x02\x00\x0f\x00" + b"\x02\x00\x17\x00\x00\x00\xff\xff"  # #REF!
    assert _XLS_ERROR_RECORD.search(b"xx" + boolerr) and _XLS_ERROR_RECORD.search(b"xx" + formula)
    assert not _XLS_ERROR_RECORD.search(b"\x05\x02\x08\x00" + b"\x01\x00\x02\x00\x0f\x00" + b"\x01\x00")  # TRUE
    assert error_cells(SPREADSHEETS[".xls"](), ".xls", "Customers") == {(1, 0): "#DIV/0!", (2, 1): "#VALUE!",
                                                                         (2, 2): "#N/A", (3, 2): "#REF!"}
