"""A spreadsheet that would need more memory than the computer has is refused with a clear message (0.15.2).

The spreadsheet reader reserves memory for the whole rectangle from the first to the last value, 32 bytes per
cell. One stray value in the last cell of a sheet (XFD1048576) made it ask for 512 GB and the whole app process
was aborted. Such a sheet is now refused before the reader starts, saying where the values reach and what to do.
Nothing that fits in memory is refused: formatting doesn't count, ordinary files read as before, and any surprise
while checking skips the check.

The two "bomb" files are read in a separate Python process: if the check ever stopped working, that process
would abort, not the test run.
"""

import io
import re
import subprocess
import sys
import zipfile
from pathlib import Path

import openpyxl
import pytest
from openpyxl.styles import Font

from converter import sheetsize
from converter.reader import ReadError, read_file
from converter.sheetsize import SheetTooBig, check_sheet_size, col_letters

ROOT = Path(__file__).resolve().parent.parent


def xlsx(cells: dict, sheets: tuple = ("Sheet1",), styled: tuple = ()) -> bytes:
    """An .xlsx with `cells` ({"A1": value}) on the first sheet; further sheets get a small table."""
    wb = openpyxl.Workbook()
    for i, name in enumerate(sheets):
        ws = wb.active if i == 0 else wb.create_sheet(name)
        ws.title = name
        if i == 0:
            for ref, v in cells.items():
                ws[ref] = v
            for ref in styled:
                ws[ref].font = Font(bold=True)  # formatting only, no value
        else:
            ws.append(["Phone"])
            ws.append(["5165550100"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def ods(rows_xml: str) -> bytes:
    """A minimal .ods whose only table holds `rows_xml` (table:table-row elements)."""
    content = ('<?xml version="1.0" encoding="UTF-8"?><office:document-content '
               'xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" '
               'xmlns:table="urn:oasis:names:tc:opendocument:xmlns:table:1.0" '
               'xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0" office:version="1.2"><office:body>'
               '<office:spreadsheet><table:table table:name="Sheet1">' + rows_xml +
               '</table:table></office:spreadsheet></office:body></office:document-content>')
    manifest = ('<?xml version="1.0" encoding="UTF-8"?><manifest:manifest '
                'xmlns:manifest="urn:oasis:names:tc:opendocument:xmlns:manifest:1.0"><manifest:file-entry '
                'manifest:full-path="/" manifest:media-type="application/vnd.oasis.opendocument.spreadsheet"/>'
                '<manifest:file-entry manifest:full-path="content.xml" manifest:media-type="text/xml"/>'
                '</manifest:manifest>')
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr(zipfile.ZipInfo("mimetype"), "application/vnd.oasis.opendocument.spreadsheet")
        z.writestr("content.xml", content)
        z.writestr("META-INF/manifest.xml", manifest)
    return buf.getvalue()


def cell(value: str, repeat: int = 1) -> str:
    rep = f' table:number-columns-repeated="{repeat}"' if repeat > 1 else ""
    return f'<table:table-cell office:value-type="string"{rep}><text:p>{value}</text:p></table:table-cell>'


def read_in_subprocess(tmp_path, name: str, data: bytes) -> subprocess.CompletedProcess:
    """read_file in a separate Python process; prints 'OK' or 'ReadError: <message>'."""
    p = tmp_path / name
    p.write_bytes(data)
    code = ("import sys; sys.path.insert(0, sys.argv[1]); from converter.reader import read_file, ReadError\n"
            "try:\n    read_file(sys.argv[2], open(sys.argv[2], 'rb').read()); print('OK')\n"
            "except ReadError as e:\n    print('ReadError:', e)\n")
    return subprocess.run([sys.executable, "-c", code, str(ROOT), str(p)], capture_output=True, text=True,
                          timeout=120)


def test_stray_value_in_last_cell_is_refused_not_a_crash(tmp_path):
    """The 5 KB file that used to abort the app gets a message naming the cell and Ctrl+End."""
    res = read_in_subprocess(tmp_path, "stray.xlsx", xlsx({"A1": "Phone", "A2": "5165550100", "XFD1048576": "x"}))
    assert res.returncode == 0, res.stderr[-500:]
    assert res.stdout.startswith("ReadError: This sheet is too big to read on this computer: its values reach "
                                 "cell XFD1048576, a block of 1,048,576 rows x 16,384 columns")
    assert "Ctrl+End" in res.stdout


def test_repeated_rows_ods_is_refused(tmp_path):
    """An .ods of a few hundred bytes that repeats a row of values a million times."""
    data = ods("<table:table-row>" + cell("Phone") + "</table:table-row>"
               '<table:table-row table:number-rows-repeated="1000000">' + cell("5165550100", 50) +
               "</table:table-row>")
    assert len(data) < 1500
    res = read_in_subprocess(tmp_path, "bomb.ods", data)
    assert res.returncode == 0, res.stderr[-500:]
    assert res.stdout.startswith("ReadError: This sheet is too big to read on this computer: its values reach "
                                 "row 1,000,001")


def test_formatting_far_away_is_not_a_reason_to_refuse():
    """Formatted but empty cells (Excel keeps them in the file) are ignored, as the reader ignores them."""
    rr = read_file("f.xlsx", xlsx({"A1": "Phone", "A2": "5165550100"}, styled=("XFD1048576", "C900000")))
    assert rr.df["Phone"].tolist() == ["5165550100"]


def test_ordinary_files_read_as_before():
    rr = read_file("f.xlsx", xlsx({"A1": "Phone", "B1": "Last Visit", "A2": "5165550100", "B2": "10/01/2026"}))
    assert list(rr.df.columns) == ["Phone", "Last Visit"] and len(rr.df) == 1


def test_only_the_sheet_being_read_is_checked():
    """A stray value on another sheet doesn't stop the first sheet; choosing that sheet is refused."""
    data = xlsx({"A1": "Phone", "A2": "5165550100", "XFD1048576": "x"}, sheets=("Stray", "Customers"))
    assert read_file("f.xlsx", data, sheet="Customers").df["Phone"].tolist() == ["5165550100"]
    with pytest.raises(SheetTooBig):
        check_sheet_size(data, ".xlsx", "Stray")
    check_sheet_size(data, ".xlsx", "Customers")


def test_limit_follows_this_computers_memory():
    """Refused only above the budget (80% of memory): the same sheet passes with a bigger budget."""
    data = xlsx({"A1": "Phone", "A2": "5165550100", "Z200000": "x"})  # 200,000 rows x 26 columns
    with pytest.raises(SheetTooBig, match=r"reach row 200,000 \(cell Z200000\)|reach cell Z200000"):
        check_sheet_size(data, ".xlsx", None, budget=100 << 20)
    check_sheet_size(data, ".xlsx", None, budget=2 << 30)
    assert sheetsize.memory_total() > 0
    assert sheetsize.memory_budget() == int(sheetsize.memory_total() * 0.8)


def test_cells_written_in_another_attribute_order_are_still_found():
    """Writers other than Excel may put the position after other attributes (<c s="1" r="XFD1048576">)."""
    data = xlsx({"A1": "Phone", "A2": "5165550100", "XFD1048576": "x"})
    src = zipfile.ZipFile(io.BytesIO(data))
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as out:
        for item in src.infolist():
            body = src.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                body = re.sub(rb'<c r="([A-Z]+[0-9]+)"([^>]*)>', rb'<c\2 r="\1">', body)
                assert b'<c r="' not in body
            out.writestr(item, body)
    with pytest.raises(SheetTooBig, match="XFD1048576"):
        check_sheet_size(buf.getvalue(), ".xlsx", None)


def test_libreoffice_style_empty_rest_of_sheet_is_fine():
    """LibreOffice writes the unused rest of a sheet as huge runs of empty rows and cells: not values."""
    data = ods("<table:table-row>" + cell("Phone") + '<table:table-cell table:number-columns-repeated="16383"/>'
               "</table:table-row><table:table-row>" + cell("5165550100") +
               '<table:table-cell table:number-columns-repeated="16383"/></table:table-row>'
               '<table:table-row table:number-rows-repeated="1048574"><table:table-cell '
               'table:number-columns-repeated="16384"/></table:table-row>')
    check_sheet_size(data, ".ods", None, budget=1 << 20)
    assert read_file("f.ods", data).df["Phone"].tolist() == ["5165550100"]


def test_check_never_blocks_a_file_it_cant_understand():
    """A corrupt or unusual file skips the check; the reader reports it as before."""
    check_sheet_size(b"PK\x03\x04 not really a zip", ".xlsx", None)
    check_sheet_size(b"anything", ".xls", None)
    with pytest.raises(ReadError, match="Could not open spreadsheet"):
        read_file("broken.xlsx", b"PK\x03\x04 not really a zip")


def test_a_file_that_passed_is_not_scanned_again(monkeypatch):
    """Re-read, refresh and the background import reuse the first check (same bytes, sheet and budget)."""
    data = xlsx({"A1": "Phone", "A2": "5165550101"})
    calls = []
    real = sheetsize._xlsx_extent
    monkeypatch.setattr(sheetsize, "_xlsx_extent", lambda *a: calls.append(1) or real(*a))
    check_sheet_size(data, ".xlsx", None)
    check_sheet_size(data, ".xlsx", None)
    assert len(calls) == 1


@pytest.mark.parametrize("n,letters", [(1, "A"), (26, "Z"), (27, "AA"), (702, "ZZ"), (703, "AAA"), (16384, "XFD")])
def test_column_letters(n, letters):
    assert col_letters(n) == letters
