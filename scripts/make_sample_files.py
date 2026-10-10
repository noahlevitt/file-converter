"""Generate deliberately messy sample files for testing (python scripts/make_sample_files.py).

Covers: title rows above the header, Windows-1252 encoding, semicolon/tab delimiters, no header row,
Excel numeric phones and ZIPs, Excel date cells, day-first dates, full names and full addresses in
one column, duplicates (same phone + same date) and non-duplicates (same phone, different date),
blank rows, and invalid values.

Usage::

    python scripts/make_sample_files.py              # writes into sample_data/ next to scripts/
    python scripts/make_sample_files.py /tmp/samples # writes into the given folder (created if missing)

Existing files with the same names are overwritten. The test suite calls :func:`make_all` with a temporary folder
(``tests/conftest.py``), so tests never depend on the committed copies in ``sample_data/``. Each writer function
takes the output folder (``OUT``, which must already exist) and writes exactly one file; the content is fixed
(no randomness), so regenerating gives the same data (the .xlsx files differ only in their internal timestamps).

Files produced (details and the edge cases of every row are in each function's docstring):

==========================  ====================================================================================
pos_guest_export.csv        POS export: Windows-1252, CRLF, title rows above the header, full names/addresses,
                            two phone columns, duplicate and non-duplicate rows, a blank row, invalid values.
loyalty_members.xlsx        Excel workbook whose data is on the 2nd sheet; numeric phone/ZIP cells, an Excel
                            serial date, an exact duplicate, an empty row, an invalid ZIP.
eu_style_semicolon.csv      Semicolon-delimited, non-US headers, day-first dates.
no_header.tsv               Tab-delimited with no header row at all.
restaurant_crm_export.csv   CRM export shaped for the real import format: Opt-In column, activity columns,
                            mixed date formats, birthdays without a year, duplicates, missing/short phones.
loyalty_activity.xlsx       Loyalty export for the real format: title row, Excel date-time cells, numeric and
                            float phones, DOB as Excel date or text without a year, an "N/A" phone.
==========================  ====================================================================================

The first four files exercise the generic features (names, addresses, ZIP/state) with the example contact
schema in ``tests/fixtures/example_contact_schema.yaml``; the last two match ``schemas/target_schema.yaml``.
"""

from __future__ import annotations

import datetime as dt
import sys
from pathlib import Path

import openpyxl

# Default output folder: <project>/sample_data (the committed sample files).
DEFAULT_OUT = Path(__file__).resolve().parent.parent / "sample_data"


def pos_export_csv(OUT: Path) -> None:
    """Write ``pos_guest_export.csv``: a point-of-sale guest export with typical real-world mess.

    File-level edge cases:
        * Encoded as Windows-1252 (what Excel's "Save as CSV" produces on US Windows), so "José Martínez" is not
          valid UTF-8 — the reader must detect cp1252.
        * CRLF line endings.
        * Three junk rows (report title, "Generated …" line, an all-commas row) above the real header on file
          row 4 — header detection must skip them; data starts on file row 5.
        * One full-name column (``Guest Name``) and one single-line address column (``Address``) that must be
          split with the name_part / address_part transforms; two phone columns (``Mobile #``, ``Home Phone``).

    Expected result with the example contact schema (see ``tests/test_engine.py``): 9 non-empty rows, 1 blank
    row skipped, 2 rejected, 1 duplicate removed, 7 rows imported (one extra row for Smith's second number).
    Each row's purpose is commented below.
    """
    rows = [
        # Title rows above the header (header detection must skip them).
        "Daily Guest Export,,,,,,",
        "Generated 10/08/2026 11:42 PM,,,,,,",
        ",,,,,,",
        "Guest Name,Mobile #,Home Phone,E-Mail,Birthday,Join Date,Address",
        # Accented name (cp1252 bytes), "(516) 555-0100" phone, upper-case e-mail, address with an "Apt 4B" unit.
        'José Martínez,(516) 555-0100,,JOSE.M@GMAIL.COM,03/14/1988,10/01/2026,"12 Main St Apt 4B, Hicksville, NY 11801"',
        # Quoted "Last, First" name; dotted phone plus a 2nd number in Home Phone (-> two output rows); invalid
        # e-mail without a TLD (optional -> blanked); single-digit month/day birthday.
        '"SMITH, JOHN",516.555.0142,516-555-0199,jsmith@yahoo,7/4/1975,10/02/2026,"455 Old Country Rd, Westbury, NY 11590"',
        # Apostrophe in the name; no mobile, only a bare 10-digit home phone; no birthday; state spelled "New York".
        'Mary O\'Neil,,9175550123,mary@oneil.com,,10/02/2026,"9 Ocean Ave, Long Beach, New York 11561"',
        # Exact copy of the first José row -> removed as a duplicate.
        'José Martínez,(516) 555-0100,,JOSE.M@GMAIL.COM,03/14/1988,10/01/2026,"12 Main St Apt 4B, Hicksville, NY 11801"',
        # Same person and phone (formatted differently) but a different Join Date -> NOT a duplicate; both import.
        'José Martínez,516-555-0100,,jose.m@gmail.com,03/14/1988,10/05/2026,"12 Main St Apt 4B, Hicksville, NY 11801"',
        # Blank row (only delimiters) -> skipped as empty, not counted as a data row.
        ",,,,,,",
        # 7-digit phone with no area code -> invalid required phone -> row rejected.
        'Bad Phone Guy,555-1234,,bad@example.com,01/01/1990,10/03/2026,"1 Fake St, Nowhere, NY 11111"',
        # Title + suffix in the name; "+1" prefix and an "x22" extension; ISO birthday; "Oct 4 2026" text date;
        # "Suite 200" goes to address line 2.
        'Dr. Priya Patel MD,+1 (631) 555-0101 x22,,PRIYA@PATELMD.COM,1990-02-28,Oct 4 2026,"77 Jericho Tpke Suite 200, Jericho, NY 11753"',
        # Bare 10-digit phone, no e-mail, impossible birthday 13/45/1999 (blanked), empty trailing address cell.
        'Kevin Lee,6315550188,,,13/45/1999,10/06/2026,',
        # No phone at all -> required phone missing -> row rejected.
        'No Phone Person,,,nophone@example.com,,10/06/2026,"5 Elm St, Freeport, NY 11520"',
    ]
    # CRLF line endings and Windows-1252 bytes, like a CSV saved by Excel on Windows.
    (OUT / "pos_guest_export.csv").write_bytes(("\r\n".join(rows) + "\r\n").encode("cp1252"))


def loyalty_xlsx(OUT: Path) -> None:
    """Write ``loyalty_members.xlsx``: a loyalty-program workbook whose data is NOT on the first sheet.

    Edge cases:
        * Sheet 1 "Summary" holds a small report; the members are on sheet 2 "Members" (sheet selection).
        * ALL-CAPS headers; names in upper case ("JOHN MCDONALD") and lower case ("ana garcia") to normalize.
        * Phones as Excel numbers (int 5165550100, float 2125550198.0) and as text ("201-555-0177") — the reader
          must not produce "2125550198.0" or scientific notation.
        * ZIPs as numbers: New Jersey 7030 lost its leading zero and must become "07030"; "07302" kept as text.
        * SIGNUP as real Excel date cells, plus one raw Excel serial number (46270 = 09/05/2026).
        * States written as "New York", "new jersey", "NJ", "NY", "N.Y.".
        * An exact duplicate row (Sam Rivera twice); Sam also shares John McDonald's phone with different
          details (not a duplicate).
        * A completely empty row, and an invalid ZIP "ABCDE" (optional -> blanked, the row still imports).

    Expected with the example contact schema (``tests/test_engine.py``): 1 duplicate, 0 rejected, 1 value blanked.
    """
    wb = openpyxl.Workbook()
    # First sheet: a summary that is not the data, so the user (or --sheet) must pick "Members".
    ws0 = wb.active
    ws0.title = "Summary"
    ws0.append(["Loyalty Program Summary"])
    ws0.append(["Members", 6])
    # Second sheet: the actual member list.
    ws = wb.create_sheet("Members")
    ws.append(["FIRST NAME", "LAST NAME", "PHONE NUMBER", "EMAIL ADDRESS", "SIGNUP", "CITY", "STATE", "ZIP CODE",
               "POINTS"])
    ws.append(["JOHN", "MCDONALD", 5165550100, "john@mcd.com", dt.datetime(2026, 9, 1), "HICKSVILLE", "New York", 11801, 120])
    ws.append(["ana", "garcia", 2015550111, "ana@garcia.net", dt.datetime(2026, 9, 2), "hoboken", "new jersey", 7030, 40])
    ws.append(["Li", "Wei", "201-555-0177", "LI.WEI@EXAMPLE.ORG", dt.datetime(2026, 9, 2), "Jersey City", "NJ", "07302", 0])
    # Sam Rivera twice, identical -> the second is a duplicate (same phone as John McDonald, but other details).
    ws.append(["Sam", "Rivera", 5165550100, "sam@rivera.com", dt.datetime(2026, 9, 3), "Hicksville", "NY", 11801, 15])
    ws.append(["Sam", "Rivera", 5165550100, "sam@rivera.com", dt.datetime(2026, 9, 3), "Hicksville", "NY", 11801, 15])
    # Float phone, signup as a bare Excel serial number (46270 -> 09/05/2026), state "N.Y.".
    ws.append(["Tara", "Mor", 2125550198.0, "tara@example.com", 46270, "New York", "N.Y.", 10001, 300])
    # Completely empty row -> skipped.
    ws.append([None, None, None, None, None, None, None, None, None])
    # Invalid optional ZIP -> blanked; the row still imports.
    ws.append(["Bad", "Zip", 7185550145, "bad.zip@example.com", dt.datetime(2026, 9, 4), "Brooklyn", "NY", "ABCDE", 5])
    wb.save(OUT / "loyalty_members.xlsx")


def eu_style_csv(OUT: Path) -> None:
    """Write ``eu_style_semicolon.csv``: a European-style export (UTF-8, LF line endings).

    Edge cases:
        * Semicolon delimiter (delimiter sniffing).
        * Non-US header wording: Name / Surname / Cell / Registered / Postcode.
        * Day-first dates: 15/03/2026 and 28/02/2026 can only be day-first, which tells the column-level date
          detection to read the ambiguous 01/04/2026 as 1 April (not January 4). Expected output:
          03/15/2026, 04/01/2026, 02/28/2026.
        * Phones written three ways: "+1 516 555 0123", "516 555 0124", "516-555-0125".
    """
    rows = [
        "Name;Surname;Cell;Email;Registered;Postcode",
        "Giulia;Rossi;+1 516 555 0123;giulia@rossi.it;15/03/2026;11801",
        "Marco;Bianchi;516 555 0124;marco@bianchi.it;01/04/2026;11590",
        "Elena;Verdi;516-555-0125;elena@verdi.it;28/02/2026;11753",
    ]
    (OUT / "eu_style_semicolon.csv").write_text("\n".join(rows) + "\n", encoding="utf-8")


def no_header_tsv(OUT: Path) -> None:
    """Write ``no_header.tsv``: three tab-separated rows (phone, e-mail, ISO date) with NO header row.

    The reader must detect that the first row is data, name the columns "Column 1", "Column 2", "Column 3", and
    the rules must map the phone and e-mail columns from their values alone. The date column stays unmapped:
    without a header, date fields need a clear header match.
    """
    rows = [
        "5165550131\tamy@example.com\t2026-08-01",
        "5165550132\tbob@example.com\t2026-08-02",
        "5165550133\tcat@example.com\t2026-08-03",
    ]
    (OUT / "no_header.tsv").write_text("\n".join(rows) + "\n", encoding="utf-8")


# ---------------------------------------------------------------------------
# Samples shaped like real source files for the generic import format
# ---------------------------------------------------------------------------

def restaurant_crm_csv(OUT: Path) -> None:
    """CRM export: full names, an Opt-In column that must NOT reach `consent`, mixed date formats,
    birthdays without a year, an exact duplicate row and a same-phone/different-visit row.

    Written as UTF-8 with LF line endings and a single header row. Shaped for the real import format
    (``schemas/target_schema.yaml``); the expected mapping (``tests/test_real_schema.py``) is:
    Mobile -> phone_number, Last Visit -> transactions_last, Last Call -> calls_last, Last Text Received ->
    in_sms_last, Last Campaign Sent -> last_contacted, Last Reward Redeemed -> redemptions_last, Birthday ->
    birthday, Anniversary -> anniversary. ``Opt-In Status`` and ``Created`` must stay unused (consent is always
    blank in this format), and nothing feeds spin_last.

    Edge cases by value:
        * Activity date-times in several styles: "10/05/2026 7:42 PM", "2026-10-03 18:15:00", "Oct 2 2026 1:05pm"
          and the MySQL zero date "0000-00-00 00:00:00" (-> blank).
        * Birthdays/anniversaries without a year ("09/12", "Mar 3", "September 12th", "12/31" -> year 1600),
          with a year ("1985-07-04", "06/20/2015") and impossible ("02/30").
        * Extra unmapped columns: Last Order Total ("$38.50"), Total Visits.

    Expected: 5 rows imported, 2 rejected, 1 duplicate (see the row comments below).
    """
    header = ("Customer Name,Mobile,Email,Opt-In Status,Last Visit,Last Order Total,Total Visits,Last Call,"
              "Last Text Received,Last Campaign Sent,Last Reward Redeemed,Birthday,Anniversary,Created")
    rows = [
        # Every activity column filled; birthday without a year ("09/12" -> 09/12/1600).
        'Maria Lopez,(516) 555-0100,maria@example.com,Yes,10/05/2026 7:42 PM,$38.50,12,10/01/2026 12:15 PM,'
        '10/04/2026 6:01 PM,10/03/2026 11:00 AM,09/28/2026 8:30 PM,09/12,06/20/2015,2024-01-15',
        # "+1" prefix; ISO date-times with seconds; Opt-In "No"; birthday "Mar 3" (-> 03/03/1600).
        'Tom Becker,+1 631-555-0101,tom@example.com,No,2026-10-03 18:15:00,$12.00,3,,,2026-10-03 11:00:00,,'
        'Mar 3,,2025-05-02',
        # Dotted phone; Opt-In "true"; "Oct 2 2026 1:05pm" style; birthday with a year.
        'Aisha Khan,917.555.0123,aisha@example.com,true,Oct 2 2026 1:05pm,$64.20,25,,10/02/2026 1:30 PM,'
        '10/03/2026 11:00 AM,10/02/2026 1:10 PM,1985-07-04,,2023-11-30',
        # Exact copy of the first Maria row -> removed as a duplicate.
        'Maria Lopez,(516) 555-0100,maria@example.com,Yes,10/05/2026 7:42 PM,$38.50,12,10/01/2026 12:15 PM,'
        '10/04/2026 6:01 PM,10/03/2026 11:00 AM,09/28/2026 8:30 PM,09/12,06/20/2015,2024-01-15',
        # Same phone (formatted differently) but a later Last Visit -> NOT a duplicate; imported as its own row.
        'Maria Lopez,516-555-0100,maria@example.com,Yes,10/07/2026 8:10 PM,$22.00,13,10/01/2026 12:15 PM,'
        '10/04/2026 6:01 PM,10/03/2026 11:00 AM,09/28/2026 8:30 PM,09/12,06/20/2015,2024-01-15',
        # No phone -> required phone_number missing -> rejected.
        'No Phone Guest,,nophone@example.com,Yes,10/06/2026 9:00 PM,$15.00,1,,,,,,,2026-10-06',
        # 7-digit phone -> invalid -> rejected (its impossible birthday "02/30" would be invalid too).
        'Short Number,555-0199,short@example.com,,10/06/2026 9:15 PM,$9.99,1,,,,,02/30,,2026-10-06',
        # Zero date (-> blank), "$0.00", "September 12th" (-> 09/12/1600), anniversary "12/31" (-> 12/31/1600).
        'Daniel Cho,6315550188,,No,0000-00-00 00:00:00,$0.00,0,,,,,September 12th,12/31,2026-10-07',
    ]
    (OUT / "restaurant_crm_export.csv").write_text(header + "\n" + "\n".join(rows) + "\n", encoding="utf-8")


def loyalty_activity_xlsx(OUT: Path) -> None:
    """Loyalty platform export: Excel date-time cells, numeric phones, DOB as Excel dates or text.

    Shaped for the real import format. Edge cases:
        * A report title row and an empty row above the header -> header detected on file row 3.
        * Phones as an Excel int, an Excel float (9175550123.0) and the text "N/A" (a null token -> phone
          missing -> that row is rejected).
        * Last Spin / Last Redeemed as real Excel date-time cells (-> spin_last / redemptions_last), some empty.
        * DOB as an Excel date (1990-02-28 -> 02/28/1990), as text without a year ("12-Sep" -> 09/12/1600) and as
          a leap day without a year ("02/29" -> 02/29/1600; 1600 is a leap year, so it stays valid).
        * Points: an extra column that is not part of the format.

    Expected (``tests/test_real_schema.py``): 3 rows imported, 1 rejected.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Activity"
    # Title row + empty row above the header.
    ws.append(["Loyalty Activity Report"])
    ws.append([])
    ws.append(["Phone", "Last Spin", "Last Redeemed", "DOB", "Points"])
    ws.append([5165550100, dt.datetime(2026, 10, 4, 18, 30), dt.datetime(2026, 9, 30, 12, 5), dt.datetime(1990, 2, 28), 140])
    ws.append([2015550111, dt.datetime(2026, 10, 1, 9, 5), None, "12-Sep", 20])
    ws.append([9175550123.0, None, dt.datetime(2026, 10, 2, 13, 10), "02/29", 75])
    # "N/A" phone -> treated as empty -> required phone missing -> rejected.
    ws.append(["N/A", dt.datetime(2026, 10, 3, 20, 0), None, None, 0])
    wb.save(OUT / "loyalty_activity.xlsx")


def make_all(out: Path = DEFAULT_OUT) -> Path:
    """Write all six sample files into ``out`` (created if needed, existing files overwritten); return ``out``.

    Used by the ``__main__`` block below and by the ``samples`` fixture in ``tests/conftest.py``.
    """
    out.mkdir(parents=True, exist_ok=True)
    for fn in (pos_export_csv, loyalty_xlsx, eu_style_csv, no_header_tsv, restaurant_crm_csv, loyalty_activity_xlsx):
        fn(out)
    return out


# Optional first argument: output folder (default: sample_data/).
if __name__ == "__main__":
    target = make_all(Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_OUT)
    print(f"Sample files written to {target}")
